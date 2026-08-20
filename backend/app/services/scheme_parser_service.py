import re
import io
from typing import List, Dict
import pdfplumber
from docx import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


def parse_scheme_file(file_content: bytes, filename: str) -> List[Dict]:
    """
    Parses a PDF or Word marking scheme document.
    Returns a list of question objects:
    { question_number, label_variants, max_marks, expected_answer, sub_questions }
    """
    if filename.lower().endswith(".pdf"):
        text = _extract_pdf_text(file_content)
    elif filename.lower().endswith((".docx", ".doc")):
        text = _extract_docx_text(file_content)
    else:
        raise ValueError("Unsupported file format. Use PDF or DOCX.")

    return _parse_scheme_text(text)


def _extract_pdf_text(content: bytes) -> str:
    text_parts = []
    with pdfplumber.open(io.BytesIO(content)) as pdf:
        for page in pdf.pages:
            # layout=True preserves reading order better for complex layouts
            # (e.g. multi-column or equations mixed with text).
            text = page.extract_text(layout=True) or page.extract_text() or ""
            text_parts.append(text)
    return "\n".join(text_parts)


_W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
_M = 'http://schemas.openxmlformats.org/officeDocument/2006/math'

# Common Unicode math characters that appear in <m:t> and their LaTeX equivalents
_UNICODE_TO_LATEX = {
    '−': '-', '×': r'\times ', '÷': r'\div ', '±': r'\pm ', '∓': r'\mp ',
    '≤': r'\leq ', '≥': r'\geq ', '≠': r'\neq ', '≈': r'\approx ',
    '∞': r'\infty ', '∈': r'\in ', '∉': r'\notin ', '⊂': r'\subset ',
    '∪': r'\cup ', '∩': r'\cap ', '∅': r'\emptyset ',
    'α': r'\alpha ', 'β': r'\beta ', 'γ': r'\gamma ', 'δ': r'\delta ',
    'ε': r'\epsilon ', 'θ': r'\theta ', 'λ': r'\lambda ', 'μ': r'\mu ',
    'π': r'\pi ', 'σ': r'\sigma ', 'φ': r'\phi ', 'ω': r'\omega ',
    'Δ': r'\Delta ', 'Σ': r'\Sigma ', 'Π': r'\Pi ', 'Ω': r'\Omega ',
}

_NARY_OP_MAP = {
    '∫': r'\int', '∬': r'\iint', '∭': r'\iiint', '∮': r'\oint',
    '∑': r'\sum', '∏': r'\prod',
}


def _omml_to_latex(elem) -> str:
    """
    Recursively converts an OMML element tree to a LaTeX string.
    Handles the most common math constructs produced by Word's equation editor.
    """
    def tag(e):
        return e.tag.split('}')[1] if '}' in e.tag else e.tag

    def ch(e):
        """Concatenate LaTeX of all children."""
        return ''.join(_omml_to_latex(c) for c in e)

    def child(e, name):
        """LaTeX of a named child element."""
        c = e.find(f'{{{_M}}}{name}')
        return ch(c) if c is not None else ''

    t = tag(elem)

    if t == 't':
        text = elem.text or ''
        for uni, lat in _UNICODE_TO_LATEX.items():
            text = text.replace(uni, lat)
        return text

    if t in ('r', 'oMath', 'oMathPara', 'e', 'num', 'den',
             'sup', 'sub', 'deg', 'lim', 'fName'):
        return ch(elem)

    if t == 'f':                      # fraction
        return rf'\frac{{{child(elem, "num")}}}{{{child(elem, "den")}}}'

    if t == 'sSup':                   # superscript
        return f'{child(elem, "e")}^{{{child(elem, "sup")}}}'

    if t == 'sSub':                   # subscript
        return f'{child(elem, "e")}_{{{child(elem, "sub")}}}'

    if t == 'sSubSup':                # sub+superscript
        return (f'{child(elem, "e")}'
                f'_{{{child(elem, "sub")}}}'
                f'^{{{child(elem, "sup")}}}')

    if t == 'rad':                    # radical / sqrt
        deg = child(elem, 'deg').strip()
        inner = child(elem, 'e')
        return rf'\sqrt[{deg}]{{{inner}}}' if deg else rf'\sqrt{{{inner}}}'

    if t == 'nary':                   # integral, sum, product, etc.
        pr = elem.find(f'{{{_M}}}naryPr')
        chr_e = pr.find(f'{{{_M}}}chr') if pr is not None else None
        op_char = chr_e.get(f'{{{_M}}}val', '∫') if chr_e is not None else '∫'
        op = _NARY_OP_MAP.get(op_char, rf'\operatorname{{{op_char}}}')
        sub_l = child(elem, 'sub').strip()
        sup_l = child(elem, 'sup').strip()
        body = child(elem, 'e')
        result = op
        if sub_l:
            result += f'_{{{sub_l}}}'
        if sup_l:
            result += f'^{{{sup_l}}}'
        return result + f' {body}'

    if t == 'd':                      # delimiters: ( ) [ ] | |
        pr = elem.find(f'{{{_M}}}dPr')
        val = f'{{{_M}}}val'
        beg = pr.find(f'{{{_M}}}begChr').get(val, '(') if pr is not None and pr.find(f'{{{_M}}}begChr') is not None else '('
        end = pr.find(f'{{{_M}}}endChr').get(val, ')') if pr is not None and pr.find(f'{{{_M}}}endChr') is not None else ')'
        return f'{beg}{child(elem, "e")}{end}'

    if t == 'func':                   # named function e.g. sin, cos
        return f'{child(elem, "fName")}({child(elem, "e")})'

    if t == 'eqArr':                  # equation array → aligned
        rows = [ch(row) for row in elem.findall(f'{{{_M}}}e')]
        return r'\begin{aligned}' + r'\\'.join(rows) + r'\end{aligned}'

    if t == 'm':                      # matrix
        rows = elem.findall(f'{{{_M}}}mr')
        row_strs = [' & '.join(ch(cell) for cell in row.findall(f'{{{_M}}}e'))
                    for row in rows]
        return r'\begin{pmatrix}' + r'\\'.join(row_strs) + r'\end{pmatrix}'

    # Unknown element — just recurse into children
    return ch(elem)


def _iter_block_items(doc: DocxDocument):
    """Yields Paragraph/Table objects from the document body in the order
    they actually appear — doc.paragraphs and doc.tables each only give one
    type and lose the interleaving, which matters here since a table (e.g.
    an answer given as a comparison table) can sit between two questions."""
    for child in doc.element.body.iterchildren():
        if child.tag == qn('w:p'):
            yield Paragraph(child, doc)
        elif child.tag == qn('w:tbl'):
            yield Table(child, doc)


def _paragraph_text(p: Paragraph) -> str:
    """
    Extracts text from a single paragraph, converting Word equation-editor
    (OMML) blocks to LaTeX so that KaTeX can render them in the frontend.

    - <m:oMathPara>  (display/centred equation) → wrapped in \\[...\\]
    - <m:oMath>       (inline equation)          → wrapped in \\(...\\)
    - <w:r>           (regular text run)          → plain text
    """
    parts = []
    for child in p._element:
        ns = child.tag.split('}')[0].lstrip('{') if '}' in child.tag else ''
        local = child.tag.split('}')[1] if '}' in child.tag else child.tag

        if ns == _W and local == 'r':
            # Regular text run
            for t_elem in child.iter(f'{{{_W}}}t'):
                parts.append(t_elem.text or '')

        elif ns == _M and local == 'oMathPara':
            # Display (centred) equation
            inner = child.find(f'{{{_M}}}oMath')
            latex = _omml_to_latex(inner if inner is not None else child)
            parts.append(f'\\[{latex}\\]')

        elif ns == _M and local == 'oMath':
            # Inline equation
            latex = _omml_to_latex(child)
            parts.append(f'\\({latex}\\)')

    return ''.join(parts)


def _escape_md_cell(text: str) -> str:
    """Escapes markdown/table-syntax-significant characters so scheme
    content (e.g. a literal '*' or '_') displays literally instead of
    being misread as formatting or breaking the table's column
    delimiters, and collapses any embedded newlines (invalid inside a
    single markdown table cell)."""
    text = re.sub(r'\s*\n\s*', ' ', text.strip())
    return re.sub(r'([\\|*_`])', r'\\\1', text)


def _extract_docx_text(content: bytes) -> str:
    """Extracts text from a DOCX file, walking paragraphs and tables in
    document order so table-based answers (e.g. a comparison table) aren't
    silently dropped — doc.paragraphs alone excludes table cell content.
    Tables are emitted as GFM markdown so they render as an actual table
    in the frontend instead of plain 'cellA | cellB' text."""
    doc = DocxDocument(io.BytesIO(content))
    lines = []
    for block in _iter_block_items(doc):
        if isinstance(block, Table):
            rows = [[_escape_md_cell(cell.text) for cell in row.cells] for row in block.rows]
            rows = [r for r in rows if any(r)]
            if rows:
                col_count = max(len(r) for r in rows)
                padded = [r + [''] * (col_count - len(r)) for r in rows]
                table_lines = ['| ' + ' | '.join(r) + ' |' for r in padded]
                table_lines.insert(1, '|' + '|'.join([' --- '] * col_count) + '|')
                lines.append('')
                lines.extend(table_lines)
                lines.append('')
        else:
            lines.append(_paragraph_text(block))
    return '\n'.join(lines)


def _parse_scheme_text(text: str) -> List[Dict]:
    """
    Parses questions from scheme text using a line-by-line approach.
    Finds header lines matching Q1 / Question 1 / etc., then extracts
    marks from the header or anywhere in the question body.

    Accepted marks formats:
      (10 marks)  [10 marks]  {10 marks}
      10 marks    Marks: 10   10/20
    """
    HEADER_RE = re.compile(
        r'^\s*(?:Q|Question|QUESTION)\s*\.?\s*(\d+[a-zA-Z]?)\b(.*)',
        re.IGNORECASE
    )
    MARKS_RE = re.compile(
        r'[\(\[\{](\d+(?:\.\d+)?)[^\S\n]*marks?[^\S\n]*[\)\]\}]'  # (10 marks) [10 marks]
        r'|marks?[^\S\n]*[:\-][^\S\n]*(\d+(?:\.\d+)?)'             # marks: 10 / marks- 10
        r'|(\d+(?:\.\d+)?)[^\S\n]*/[^\S\n]*\d+(?:\.\d+)?'         # 10/20 — take numerator
        r'|(\d+(?:\.\d+)?)[^\S\n]*marks?\b',                        # 10 marks (same line)
        re.IGNORECASE
    )

    lines = text.split('\n')

    # Collect header positions
    headers: List[tuple] = []  # (line_index, q_num, rest_of_header_line)
    for i, line in enumerate(lines):
        m = HEADER_RE.match(line)
        if m:
            headers.append((i, m.group(1).strip(), m.group(2)))

    questions = []
    for idx, (line_no, q_num, rest_of_header) in enumerate(headers):
        next_line_no = headers[idx + 1][0] if idx + 1 < len(headers) else len(lines)

        # Search for marks: header remainder first, then full block
        block = '\n'.join(lines[line_no:next_line_no])
        mm = MARKS_RE.search(rest_of_header) or MARKS_RE.search(block)
        if mm:
            raw = mm.group(1) or mm.group(2) or mm.group(3) or mm.group(4)
            max_marks = float(raw)
        else:
            max_marks = 0.0

        # The question wording is very often on the same line as the "Q1."
        # header itself (e.g. "Q1. What is Polymorphism? (10 marks)") rather
        # than on its own line below — strip the marks annotation and the
        # leftover separator punctuation, then treat it as the first line of
        # the question body so it isn't silently dropped from question_text.
        header_wording = MARKS_RE.sub('', rest_of_header)
        header_wording = re.sub(r'^[\s.:\-]+', '', header_wording).strip()

        body = '\n'.join(lines[line_no + 1:next_line_no]).strip()
        full_body = f'{header_wording}\n{body}' if header_wording else body
        question_text, expected_answer = _split_question_body(full_body)

        questions.append({
            "question_number": q_num,
            "label_variants": _generate_label_variants(q_num),
            "max_marks": max_marks,
            "question_text": question_text,
            "expected_answer": expected_answer,
            "sub_questions": []
        })

    return questions


def _split_question_body(body: str):
    """
    Splits the body text under a question header into (question_text, expected_answer).
    Looks for a line starting with 'Expected answer', 'Answer', or 'Model answer'.
    Everything before that line = question text; everything after = expected answer.
    If no such separator is found, the whole body is question_text and expected_answer is ''.
    """
    ANSWER_PREFIX = re.compile(
        r'^(?:expected\s+answer|model\s+answer|answer)[:\s]+',
        re.IGNORECASE
    )
    lines = body.split('\n')
    split_idx = None
    for i, line in enumerate(lines):
        if ANSWER_PREFIX.match(line.strip()):
            split_idx = i
            break

    if split_idx is None:
        return body, ''

    question_text = '\n'.join(lines[:split_idx]).strip()
    # Include the separator line content (after the label) in expected_answer
    sep_line = lines[split_idx].strip()
    sep_content = ANSWER_PREFIX.sub('', sep_line).strip()
    rest = '\n'.join(lines[split_idx + 1:]).strip()
    expected_answer = (sep_content + '\n' + rest).strip() if sep_content else rest
    return question_text, expected_answer


def _generate_label_variants(q_num: str) -> List[str]:
    return [
        f"Q{q_num}", f"q{q_num}", f"Q.{q_num}", f"q.{q_num}",
        f"Question {q_num}", f"question {q_num}", f"QUESTION {q_num}",
        f"Q {q_num}", f"q {q_num}"
    ]
